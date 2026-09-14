"""Elle etiketleme sayfasi: CSV -> xlsx -> CSV gidis donusu (Hafta 2, Hafta 4).

Neden xlsx? Etiketleme dosyasi CSV olarak uretiliyor ama insanin onu Excel'de
acip kaydetmesi iki sessiz hataya aciktir: Turkce Windows yerel ayarinda liste
ayraci `;` oldugu icin virgullu dosya tek kolona duser, ve Excel BOM'suz UTF-8'i
taniyamayip cp1254 ile geri yazar - review metnindeki karakterler bozulur.
Ikisi de ancak saatler suren etiketleme bittikten sonra fark edilir. xlsx'te
ayrac ve kodlama diye bir kavram yok.

ONYARGI KONTROLU - `BLIND_COLUMNS`. Iki ayri sizinti kapatiliyor:

  1. Sozcuksel vekilin karari (`kw_gift_proxy`, `trial_stratum`). Gorulurse
     etiketler vekilin hatalarini tekrarlar ve insan-vekil karsilastirmasi
     kendi kendini dogrulayan bir olcume doner.
  2. LLM'in KENDI cevabi (`purchase_type`, `confidence`, `recipient`,
     `occasion`, `evidence_span`, `val_stratum`). Hafta 4'te olculen sey
     detektorun DOGRULUGU; etiketleyen kisi modelin cevabini gorurse olculen
     sey telkine uyum olur. `evidence_span` etiketlemeyi hizlandirirdi ama
     modelin gerekcesini de gosterir - bagimsizlik hiza tercih ediliyor
     (ayni gerekce: DECISIONS 2026-08-26, "sorun dogrulukta degil bagimsizlikta").

Kolonlar `ingest` sirasinda kimlik kolonu uzerinden geri eklenir.

Kullanim:
    # Hafta 2 - tek etiketleyici, prompt gelistirme
    python -m gift_contamination.data.labelsheet --export
    python -m gift_contamination.data.labelsheet --ingest

    # Hafta 4 - uc etiketleyici, dogrulama seti
    python -m gift_contamination.data.labelsheet --validation --export --annotators 3
    python -m gift_contamination.data.labelsheet --validation --ingest --annotators 3
"""

from __future__ import annotations

import argparse
import string
from pathlib import Path

import polars as pl

from ..config import Config
from ..data.sampling import LABELS, validation_path
from ..utils.io import should_skip
from ..utils.logging import get_logger, log_output

log = get_logger("data.labelsheet")

# Etiketin NEREDEN geldigini kayda geciren damga. `ingest` disinda bir yoldan
# uretilmis (elle yazilmis, sentetik, bir betikle doldurulmus) bir dosya
# Hafta 4 olcumune giremez: `analysis.validation.load_labels` bu degeri arar.
# Gerekce deneyimle sabit - 2026-08-28'de taklit LLM ciktisi gercek gorunen bir
# Kapi 1 karari uretmisti; ayni tuzak insan etiketi tarafinda da var ve orada
# uc kisinin saatleri soz konusu.
LABEL_SOURCE_COLUMN = "label_source"
LABEL_SOURCE = "xlsx_ingest"

# Kimlik kolonu: deneme setinde `trial_id`, dogrulama setinde `val_id`.
# Kolon adini sabitlemek yerine tespit ediyoruz - iki akis ayni koddan gecsin.
ID_COLUMNS = ("trial_id", "val_id")

# Etiketleyene GOSTERILEN kolonlar (kimlik kolonu basa eklenir). Kasitli kisa.
VISIBLE_COLUMNS = ["category", "title", "text", "label", "notes"]

# Deneme akisinin (Hafta 2) kolon duzeni. `export_columns` ile ayni sey; geriye
# donuk uyumluluk icin sabit olarak da duruyor.
EXPORT_COLUMNS = ["trial_id", *VISIBLE_COLUMNS]

# Sayfaya yazilmayan, ingest'te geri eklenen kolonlar. Modul docstring'i neden
# oldugunu anlatiyor; `test_export_hides_the_bias_columns` kilitliyor.
BLIND_COLUMNS = [
    "row_id",
    # sozcuksel vekil
    "trial_stratum", "sample_frame", "kw_gift_proxy",
    # LLM'in kendi cevabi - Hafta 4'un olcecegi sey tam olarak bu
    "val_stratum", "purchase_type", "confidence", "recipient", "occasion",
    "evidence_span",
]

_WIDTHS = {
    "trial_id": 8, "val_id": 8, "category": 16,
    "title": 34, "text": 96, "label": 14, "notes": 30,
}


def export_columns(id_col: str) -> list[str]:
    return [id_col, *VISIBLE_COLUMNS]


def _id_column(df: pl.DataFrame) -> str:
    """CSV'nin kimlik kolonunu tespit eder."""
    found = [c for c in ID_COLUMNS if c in df.columns]
    if len(found) != 1:
        raise ValueError(
            f"kimlik kolonu belirsiz: {found or 'yok'}. "
            f"Tam olarak biri bulunmali: {', '.join(ID_COLUMNS)}"
        )
    return found[0]


def annotator_tags(n: int) -> list[str | None]:
    """1 -> [None] (tek dosya), 3 -> ['A','B','C'].

    Tek etiketleyicide dosya adi degismiyor: Hafta 2 akisi aynen calisiyor.
    """
    if n < 1 or n > len(string.ascii_uppercase):
        raise ValueError(f"annotators 1..26 arasinda olmali, {n} verildi")
    return [None] if n == 1 else list(string.ascii_uppercase[:n])


def resolve_tags(annotators: int = 1, tags: list[str] | None = None) -> list[str | None]:
    """Hangi sayfalar okunacak: acik etiket listesi varsa o, yoksa sayidan.

    Acik liste GEREKLI cunku "kac kisi" ile "hangileri" ayni soru degil. Hafta 4
    uc sayfayla (A, B, C) uretildi ama yalnizca A teslim etti (2026-09-14).
    `annotators=1` o durumda `validation_500.xlsx`i (soneksiz) arardi ve A'nin
    dosyasini hic gormezdi; `tags=['A']` dosyayi ve `label_A` kolonunu korur.
    """
    if tags is None:
        return annotator_tags(annotators)
    if not tags:
        raise ValueError("etiketleyici listesi bos")
    for t in tags:
        if not (isinstance(t, str) and len(t) == 1 and t in string.ascii_uppercase):
            raise ValueError(f"etiketleyici etiketi tek buyuk harf olmali, '{t}' verildi")
    if len(set(tags)) != len(tags):
        raise ValueError(f"tekrar eden etiketleyici etiketi: {tags}")
    return list(tags)


# Rehberin istedigi not kodlari. Etiketleyici bunlari DUZ TURKCE yazabilir:
# A'nin 131 notunun hepsi "kendi çocuğu" (2026-09-14) - kod "KENDI_COCUGU"
# tam eslesme aradigi icin hepsi sifir sayiliyordu. Niyet acik, bilgi kurtarilir.
NOTE_CODES = {
    "kendicocugu": "KENDI_COCUGU",
    "emindegil": "EMIN_DEGIL",
    "emindegilim": "EMIN_DEGIL",
}
_TR_FOLD = str.maketrans("çğıöşüâîûÇĞİIÖŞÜ", "cgiosuaiucgiiosu")


def normalize_note(raw: str | None) -> str | None:
    """Not metnini koda cevirir; kod degilse metni OLDUGU GIBI birakir.

    Eslesme BUTUN not uzerinden, parca uzerinden degil: "kendi çocuğu değil"
    tam tersini soyluyor ve KENDI_COCUGU'ya donusmemeli. Harf disi her sey ve
    Turkce karakterler katlaniyor, yani "Kendi Çocuğu", "kendi_cocugu" ve
    "KENDI_COCUGU" ayni koda gidiyor.
    """
    if raw is None:
        return None
    text = str(raw).strip()
    if not text:
        return None
    katlanmis = "".join(ch for ch in text.translate(_TR_FOLD).lower() if ch.isalpha())
    return NOTE_CODES.get(katlanmis, text)


def sheet_path(csv_path: Path, tag: str | None = None) -> Path:
    stem = csv_path.stem if tag is None else f"{csv_path.stem}_{tag}"
    return csv_path.with_name(stem + ".xlsx")


def labeled_path(csv_path: Path) -> Path:
    return csv_path.with_name(csv_path.stem + "_labeled.csv")


def default_csv(cfg: Config) -> Path:
    """En son uretilen prompt deneme CSV'si."""
    found = sorted(
        p
        for p in cfg.path("human").glob("prompt_trial_*.csv")
        if not p.stem.endswith("_labeled")
    )
    if not found:
        raise FileNotFoundError(
            "prompt_trial_*.csv bulunamadi. Once uretin:\n"
            "  python -m gift_contamination.data.sampling --trial 200 --category all"
        )
    return found[-1]


def validation_csv(cfg: Config) -> Path:
    """Hafta 4'un dogrulama CSV'si. Yol sozlesmesi `sampling`de."""
    path = validation_path(cfg, int(cfg.get("validation.n")))
    if not path.exists():
        raise FileNotFoundError(
            f"{path.name} yok. Once cekin:\n"
            "  python -m gift_contamination.data.sampling --validation"
        )
    return path


# --------------------------------------------------------------------- export
def export(
    csv_path: Path, *, annotators: int = 1, tags: list[str] | None = None,
    force: bool = False,
) -> list[Path]:
    """CSV'yi acilir menulu, dondurulmus baslikli xlsx sayfalarina cevirir.

    Birden fazla etiketleyicide her dosya AYNI satirlari AYNI sirada tasir.
    Kappa hizalamasi kimlik kolonu uzerinden yapiliyor ama sirayi da sabit
    tutmak, birinin sayfayi siralayip kaydetmesi halinde farki gorunur kilar.
    """
    source = pl.read_csv(csv_path)
    id_col = _id_column(source)
    columns = export_columns(id_col)
    df = source.select(columns)

    written = []
    secilen = resolve_tags(annotators, tags)
    for tag in secilen:
        dest = sheet_path(csv_path, tag)
        if should_skip(dest, force, log):
            written.append(dest)
            continue
        _write_sheet(df, columns, dest, tag)
        written.append(dest)
        log_output(log, dest, n_rows=df.height)

    log.info("etiket kolonu acilir menulu; onyargi kolonlari sayfada YOK")
    if len(secilen) > 1:
        log.info(
            "%d ayri sayfa: her etiketleyici KENDI dosyasini doldurur, "
            "birbirininkini gormez", len(secilen),
        )
    return written


def _write_sheet(
    df: pl.DataFrame, columns: list[str], dest: Path, tag: str | None
) -> None:
    import xlsxwriter

    book = xlsxwriter.Workbook(str(dest), {"strings_to_urls": False})
    sheet = book.add_worksheet("etiketleme" if tag is None else f"etiketleme_{tag}")

    head = book.add_format({"bold": True, "bg_color": "#DDDDDD", "border": 1})
    wrap = book.add_format({"text_wrap": True, "valign": "top"})
    plain = book.add_format({"valign": "top"})
    entry = book.add_format({"valign": "top", "bg_color": "#FFF7CC", "border": 1})

    for col, name in enumerate(columns):
        sheet.write(0, col, name, head)
        fmt = wrap if name in ("title", "text") else (entry if name in ("label", "notes") else plain)
        sheet.set_column(col, col, _WIDTHS[name], fmt)

    for row, record in enumerate(df.iter_rows(named=True), start=1):
        for col, name in enumerate(columns):
            value = record[name]
            sheet.write(row, col, "" if value is None else value)

    label_col = columns.index("label")
    sheet.data_validation(
        1,
        label_col,
        df.height,
        label_col,
        {
            "validate": "list",
            "source": list(LABELS),
            "error_title": "Gecersiz etiket",
            "error_message": "Yalnizca: " + ", ".join(LABELS),
        },
    )
    sheet.freeze_panes(1, 0)
    sheet.autofilter(0, 0, df.height, len(columns) - 1)
    book.close()


# --------------------------------------------------------------------- ingest
def ingest(
    csv_path: Path, *, annotators: int = 1, tags: list[str] | None = None,
    force: bool = False,
) -> tuple[Path, dict]:
    """Doldurulmus xlsx'leri okur, dogrular ve UTF-8 CSV olarak geri yazar.

    Etiketsiz tek sayfada (Hafta 2) cikti `label` / `notes` tasir. Harfli
    sayfalarda `label_A`, `notes_A`, ... olur - tek harf de olsa (`tags=['A']`).

    `notes` KOD olarak normalize edilir (bkz. `normalize_note`); etiketleyicinin
    yazdigi ham metin `notes_raw` kolonunda degismeden durur.
    """
    source = pl.read_csv(csv_path)
    id_col = _id_column(source)
    columns = export_columns(id_col)
    secilen = resolve_tags(annotators, tags)
    etiketsiz = secilen == [None]

    reports: dict[str, dict] = {}
    merged = source.drop("label", "notes")
    for tag in secilen:
        filled = _read_sheet(csv_path, columns, id_col, tag)
        rapor = _validate(filled, source, id_col)
        normal = pl.Series(
            "notes", [normalize_note(v) for v in filled["notes"].to_list()], dtype=pl.String
        )
        rapor["notes"] = {
            "n_kendi_cocugu": int((normal == "KENDI_COCUGU").sum()),
            "n_emin_degil": int((normal == "EMIN_DEGIL").sum()),
            "n_serbest": int(
                (normal.is_not_null() & ~normal.is_in(list(set(NOTE_CODES.values())))).sum()
            ),
        }
        reports["tek" if tag is None else tag] = rapor
        suffix = "" if tag is None else f"_{tag}"
        merged = merged.join(
            filled.select(
                id_col,
                pl.col("label").alias(f"label{suffix}"),
                normal.alias(f"notes{suffix}"),
                pl.col("notes").alias(f"notes_raw{suffix}"),
            ),
            on=id_col,
            how="left",
        )

    if etiketsiz:
        merged = merged.select([*source.columns, "notes_raw"])
    merged = merged.with_columns(pl.lit(LABEL_SOURCE).alias(LABEL_SOURCE_COLUMN))

    report = reports["tek"] if etiketsiz else {"by_annotator": reports}
    dest = labeled_path(csv_path)
    if should_skip(dest, force, log):
        return dest, report
    merged.write_csv(dest)
    log_output(log, dest, n_rows=merged.height)
    return dest, report


def _read_sheet(
    csv_path: Path, columns: list[str], id_col: str, tag: str | None
) -> pl.DataFrame:
    src = sheet_path(csv_path, tag)
    if not src.exists():
        raise FileNotFoundError(f"{src} yok - once --export calistirin")

    import openpyxl

    book = openpyxl.load_workbook(src, data_only=True)
    rows = list(book.active.iter_rows(values_only=True))
    header = [str(c) if c is not None else "" for c in rows[0]]
    if header != columns:
        raise ValueError(f"basliklar degismis: {header} != {columns}")

    records = [dict(zip(header, r)) for r in rows[1:] if r[0] is not None]
    return pl.DataFrame(
        records,
        schema={
            id_col: pl.Int64,
            "category": pl.String,
            "title": pl.String,
            "text": pl.String,
            "label": pl.String,
            "notes": pl.String,
        },
    ).with_columns(pl.col("label").str.strip_chars().replace("", None))


def _validate(filled: pl.DataFrame, source: pl.DataFrame, id_col: str) -> dict:
    """Satir kaybi, tekrar ve sozluk disi etiket arar; ilerleme raporu doner."""
    missing = set(source[id_col]) - set(filled[id_col])
    extra = set(filled[id_col]) - set(source[id_col])
    if missing or extra:
        raise ValueError(
            f"{id_col} kumesi degismis - satir silinmis veya eklenmis. "
            f"eksik={sorted(missing)[:5]} fazla={sorted(extra)[:5]}"
        )
    if filled[id_col].n_unique() != filled.height:
        raise ValueError(f"tekrar eden {id_col} var - satir kopyalanmis")

    bad = filled.filter(
        pl.col("label").is_not_null() & ~pl.col("label").is_in(list(LABELS))
    ).select(id_col, "label")
    if bad.height:
        raise ValueError(
            f"{bad.height} satirda sozluk disi etiket var. Gecerli: {', '.join(LABELS)}\n{bad.head(10)}"
        )

    done = filled.filter(pl.col("label").is_not_null())
    report = {
        "n_total": filled.height,
        "n_labeled": done.height,
        "by_label": dict(done.group_by("label").len().sort("len", descending=True).iter_rows()),
    }
    if done.height and "kw_gift_proxy" in source.columns:
        # Insan ile sozcuksel vekil nerede ayrisiyor: prompt v2'nin asil girdisi,
        # Hafta 4'te de vekilin ILK bagimsiz precision'i.
        joined = done.join(source.select(id_col, "kw_gift_proxy"), on=id_col)
        gift = pl.col("label") == "gift_given"
        proxy = pl.col("kw_gift_proxy")
        report["vs_keyword_proxy"] = {
            "proxy_dogru": joined.filter(proxy & gift).height,
            "proxy_yanlis_pozitif": joined.filter(proxy & ~gift).height,
            "proxy_kacirdi": joined.filter(~proxy & gift).height,
            "ikisi_de_hayir": joined.filter(~proxy & ~gift).height,
        }
    return report


def _log_report(name: str, report: dict) -> None:
    log.info("[%s] etiketlenen: %d / %d", name, report["n_labeled"], report["n_total"])
    for label, n in report["by_label"].items():
        log.info("    %-12s %4d", label, n)
    if "vs_keyword_proxy" in report:
        log.info("    sozcuksel vekile karsi: %s", report["vs_keyword_proxy"])
    if "notes" in report:
        log.info("    notlar: %s", report["notes"])
    if report["n_labeled"] < report["n_total"]:
        log.warning(
            "    %d satir bos - etiketleme yarim",
            report["n_total"] - report["n_labeled"],
        )


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--config", default="configs/base.yaml")
    parser.add_argument("--csv", type=Path, help="Kaynak CSV (varsayilan: en son prompt_trial_*.csv)")
    parser.add_argument(
        "--validation",
        action="store_true",
        help="Hafta 4'un dogrulama setiyle calis (prompt deneme seti yerine)",
    )
    parser.add_argument("--export", action="store_true", help="xlsx etiketleme sayfasi uret")
    parser.add_argument("--ingest", action="store_true", help="Doldurulmus xlsx'i CSV'ye geri yaz")
    parser.add_argument(
        "--annotators",
        type=int,
        default=None,
        metavar="N",
        help="Kac ayri sayfa, A'dan baslayarak (dogrulamada varsayilan: validation.annotators)",
    )
    parser.add_argument(
        "--tags",
        default=None,
        metavar="A,B",
        help="Tam olarak hangi sayfalar (ör. 'A'). --annotators'i ezer.",
    )
    parser.add_argument("--force", action="store_true")
    args = parser.parse_args(argv)

    if args.export == args.ingest:
        parser.error("--export veya --ingest secin")

    cfg = Config.load(args.config)
    if args.csv:
        csv_path = args.csv
    elif args.validation:
        csv_path = validation_csv(cfg)
    else:
        csv_path = default_csv(cfg)

    annotators, tags = 1, None
    if args.tags:
        tags = [t.strip() for t in args.tags.split(",") if t.strip()]
    elif args.annotators is not None:
        annotators = args.annotators
    elif args.validation:
        # Hafta 4: config'te yazan etiketleyiciler - ör. yalnizca teslim eden A.
        tags = list(cfg.get("validation.annotators"))

    if args.export:
        export(csv_path, annotators=annotators, tags=tags, force=args.force)
        return 0

    _, report = ingest(csv_path, annotators=annotators, tags=tags, force=args.force)
    if "by_annotator" in report:
        for name, sub in report["by_annotator"].items():
            _log_report(name, sub)
    else:
        _log_report("tek", report)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
